from pathlib import Path
from typing import List, Dict, Any, Optional, Callable, Set
from google.cloud import bigquery
from google.cloud.exceptions import NotFound
from google.oauth2 import service_account

from src.config import (
    GCP_CREDENTIALS_PATH,
    DEFAULT_DATASET_ID,
    DEFAULT_TABLE_ID,
    BATCH_SIZE
)
from src.storage.db_manager import DatabaseManager


class BigQuerySyncManager:
    def __init__(
        self,
        credentials_path: Path = GCP_CREDENTIALS_PATH,
        dataset_id: str = DEFAULT_DATASET_ID,
        table_id: str = DEFAULT_TABLE_ID,
        db_manager: Optional[DatabaseManager] = None
    ):
        self.credentials_path = Path(credentials_path)
        self.dataset_id = dataset_id
        self.table_id = table_id
        self.db_manager = db_manager or DatabaseManager()
        self.client: Optional[bigquery.Client] = None
        self._table_ref = None

    def authenticate(self) -> bool:
        if not self.credentials_path.exists():
            raise FileNotFoundError(f"No se encontró el archivo: {self.credentials_path}")

        credentials = service_account.Credentials.from_service_account_file(
            str(self.credentials_path),
            scopes=["https://www.googleapis.com/auth/cloud-platform"]
        )
        self.client = bigquery.Client(credentials=credentials, project=credentials.project_id)
        self._table_ref = bigquery.DatasetReference(self.client.project, self.dataset_id).table(self.table_id)
        return True

    def get_schema(self) -> List[bigquery.SchemaField]:
        # 19 columnas: datos de negocio + mail_id al final
        return [
            bigquery.SchemaField("cuit", "STRING", mode="REQUIRED"),
            bigquery.SchemaField("razon_social", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("nombre", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("tier", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("terminal", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("mc", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("denominacion", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("calle", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("altura", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("piso", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("departamento", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("codigo_postal", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("localidad", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("provincia", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("telefono", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("vendedor", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("otros", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("fecha_email", "STRING", mode="NULLABLE"),
            bigquery.SchemaField("mail_id", "STRING", mode="NULLABLE"),  # Columna 19 para trazabilidad
        ]

    def recreate_table(self):
        if not self.client:
            self.authenticate()
        print("   🗑️ Eliminando tabla previa en BigQuery...")
        self.client.delete_table(self._table_ref, not_found_ok=True)
        print("   ✨ Creando tabla nueva con mail_id como columna 19...")
        table = bigquery.Table(self._table_ref, schema=self.get_schema())
        table.clustering_fields = ["cuit"]
        self.client.create_table(table)
        self.db_manager.mark_all_unsynced()
        print("   ✅ Tabla lista y registros locales reiniciados para sincronizar.")

    def ensure_infrastructure(self):
        if not self.client:
            self.authenticate()

        dataset_ref = bigquery.DatasetReference(self.client.project, self.dataset_id)
        try:
            self.client.get_dataset(dataset_ref)
        except NotFound:
            dataset = bigquery.Dataset(dataset_ref)
            dataset.location = "southamerica-east1"
            self.client.create_dataset(dataset, timeout=30)

        try:
            self.client.get_table(self._table_ref)
        except NotFound:
            table = bigquery.Table(self._table_ref, schema=self.get_schema())
            table.clustering_fields = ["cuit"]
            self.client.create_table(table)

    def get_existing_mail_ids(self) -> Set[str]:
        """Descarga únicamente los mail_id presentes en BigQuery (rápido y gratuito)."""
        try:
            table = self.client.get_table(self._table_ref)
            rows = self.client.list_rows(table, selected_fields=[
                bigquery.SchemaField("mail_id", "STRING")
            ])
            return {str(r["mail_id"]) for r in rows if r["mail_id"]}
        except NotFound:
            return set()

    def sync_batch(self, batch_size: int = BATCH_SIZE) -> int:
        self.ensure_infrastructure()

        rows = self.db_manager.get_pending_bq_records(limit=batch_size)
        if not rows:
            return 0

        cols = [
            "cuit", "razon_social", "nombre", "tier", "terminal", "mc",
            "denominacion", "calle", "altura", "piso", "departamento",
            "codigo_postal", "localidad", "provincia", "telefono",
            "vendedor", "otros", "fecha_email", "mail_id"
        ]

        payload = []
        record_ids = []
        for r in rows:
            record_ids.append(r["id"])
            item = {c: str(r[c] or "") for c in cols}
            payload.append(item)

        job_config = bigquery.LoadJobConfig(
            schema=self.get_schema(),
            source_format=bigquery.SourceFormat.NEWLINE_DELIMITED_JSON,
            write_disposition=bigquery.WriteDisposition.WRITE_APPEND,
        )

        load_job = self.client.load_table_from_json(payload, self._table_ref, job_config=job_config)
        load_job.result()

        if load_job.errors:
            raise RuntimeError(f"Error en BigQuery: {load_job.errors}")

        self.db_manager.mark_records_synced(record_ids)
        return len(record_ids)

    def sync_all_pending(self, progress_callback: Optional[Callable[[int, int], None]] = None) -> int:
        self.ensure_infrastructure()

        # Comprobar qué mail_ids ya existen en la nube para no duplicar jamás
        cloud_ids = self.get_existing_mail_ids()
        if cloud_ids:
            self.db_manager.mark_already_in_bq(cloud_ids)

        total_synced = 0
        while True:
            synced = self.sync_batch(batch_size=BATCH_SIZE)
            if synced == 0:
                break
            total_synced += synced
            if progress_callback:
                progress_callback(synced, total_synced)

        return total_synced


    def download_to_excel(self, destination_path: Path | str, progress_callback: Optional[Callable[[str], None]] = None) -> int:
        """Descarga toda la tabla de BigQuery directamente a un archivo .xlsx con diseño pro."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        if not self.client:
            self.authenticate()

        if progress_callback:
            progress_callback("Consultando tabla en BigQuery...")

        table = self.client.get_table(self._table_ref)
        total_rows = table.num_rows

        if progress_callback:
            progress_callback(f"Descargando {total_rows:,} registros...")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Clientes GCNET"

        # Encabezados de columnas según el esquema de BigQuery
        headers = [field.name for field in table.schema]
        ws.append(headers)

        # Diseño estético para los encabezados (Azul corporativo, texto blanco negrita)
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Inserción de filas desde BigQuery
        row_count = 0
        for row in self.client.list_rows(table):
            ws.append([str(row[h]) if row[h] is not None else "" for h in headers])
            row_count += 1
            if progress_callback and row_count % 5000 == 0:
                progress_callback(f"Procesando {row_count:,} de {total_rows:,} filas...")

        # Congelar la primera fila para facilitar navegación
        ws.freeze_panes = "A2"

        # Autoajuste del ancho de cada columna para que no se corten los textos
        for col_idx, col in enumerate(ws.columns, 1):
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        wb.save(str(destination_path))
        return row_count
